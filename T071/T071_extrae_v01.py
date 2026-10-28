#!/usr/bin/env python3
"""T071 - Inversion publica de Mexico desde 1970 y Penn World Table 10.01.

Uso (desde la carpeta T071/):
    python3 T071_extrae_v01.py

Pasos:
  (a) descarga a fuentes/ cada direccion de T071_fuentes_v01.csv cuyo archivo falte;
  (b) comprueba la huella SHA-256 de cada archivo contra T071_fuentes_v01.csv;
  (c) vuelve a leer cada cifra de T071_inversion_historica_v01.csv en su pagina y
      renglon/columna (PDF), en su hoja y celda (Excel) o en su posicion (JSON) y
      escribe T071_validacion_v01.csv;
  (d) rehace T071_pwt_mex_v01.csv y T071_pwt_leyenda_v01.csv desde fuentes/pwt1001.xlsx.

Bibliotecas:
  - Python 3.9 o posterior, biblioteca estandar: csv, decimal, difflib, functools, hashlib, json,
    pathlib, re, sys, time, unicodedata, urllib.
  - pdfplumber (probado con 0.11) para leer la capa de texto de los PDF.
  - openpyxl (probado con 3.1) para leer pwt1001.xlsx.

Las cifras marcadas con lectura = imagen se leyeron de la imagen de la pagina porque la
capa de texto (OCR del escaneo) falta o esta danada; no se pueden volver a leer por
programa y en la validacion quedan con coincide = no_aplica.
"""

import csv
import difflib
import functools
import hashlib
import json
import re
import sys
import time
import unicodedata
import urllib.request
from decimal import Decimal, InvalidOperation
from pathlib import Path

BASE = Path(__file__).resolve().parent
FUENTES = BASE / "fuentes"
CSV_FUENTES = BASE / "T071_fuentes_v01.csv"
CSV_CIFRAS = BASE / "T071_inversion_historica_v01.csv"
CSV_VALIDACION = BASE / "T071_validacion_v01.csv"
CSV_PWT = BASE / "T071_pwt_mex_v01.csv"
CSV_LEYENDA = BASE / "T071_pwt_leyenda_v01.csv"
PWT_XLSX = FUENTES / "pwt1001.xlsx"


def leer_csv(ruta):
    with open(ruta, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


# ---------------------------------------------------------------- (a) descarga

def paso_a():
    FUENTES.mkdir(exist_ok=True)
    for fila in leer_csv(CSV_FUENTES):
        destino = FUENTES / fila["archivo"]
        if destino.exists() or not fila["url"]:
            continue
        print(f"(a) descargando {fila['archivo']}")
        try:
            with urllib.request.urlopen(fila["url"], timeout=600) as r, open(destino, "wb") as f:
                while True:
                    bloque = r.read(1 << 20)
                    if not bloque:
                        break
                    f.write(bloque)
        except Exception as e:  # se informa y se sigue con los demas archivos
            print(f"    no se pudo descargar: {e}")
            if destino.exists():
                destino.unlink()
        time.sleep(1.5)


# ---------------------------------------------------------------- (b) huellas

def paso_b():
    estado = {}
    for fila in leer_csv(CSV_FUENTES):
        ruta = FUENTES / fila["archivo"]
        if not ruta.exists():
            estado[fila["archivo"]] = "falta"
        elif not fila["sha256"]:
            estado[fila["archivo"]] = "sin_huella_registrada"
        else:
            estado[fila["archivo"]] = "ok" if sha256(ruta) == fila["sha256"] else "NO_COINCIDE"
        print(f"(b) {estado[fila['archivo']]:>22}  {fila['archivo']}")
    return estado


# ---------------------------------------------------------------- (c) relectura

def normaliza(texto):
    t = unicodedata.normalize("NFKD", texto)
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", t.lower())


NUM = re.compile(r"^\(?-?[\d][\d,.]*\)?$")


def es_numero(palabra):
    return bool(NUM.match(palabra)) and any(c.isdigit() for c in palabra)


def a_decimal(texto):
    t = texto.replace(" ", "").replace(",", "")
    negativo = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        v = Decimal(t)
    except InvalidOperation:
        return None
    return -v if negativo else v


def lineas_de_pagina(pagina):
    """Agrupa las palabras de la pagina en renglones por su coordenada vertical."""
    palabras = sorted(pagina.dedupe_chars().extract_words(), key=lambda w: (w["top"], w["x0"]))
    lineas = []
    for w in palabras:
        if lineas and w["top"] - lineas[-1][0]["top"] <= 4:
            lineas[-1].append(w)
        else:
            lineas.append([w])
    return [sorted(l, key=lambda w: w["x0"]) for l in lineas]


def celdas_numericas(linea):
    """Une palabras numericas contiguas (separador de miles con espacio)."""
    celdas = []
    for w in linea:
        if not es_numero(w["text"]):
            continue
        if celdas and w["x0"] - celdas[-1]["x1"] <= 4:
            celdas[-1]["text"] += " " + w["text"]
            celdas[-1]["x1"] = w["x1"]
        else:
            celdas.append({"text": w["text"], "x0": w["x0"], "x1": w["x1"]})
    return celdas


def etiqueta(linea):
    """Texto del renglon antes de la primera cifra (o la primera palabra si es un ano)."""
    if linea and re.fullmatch(r"\d{4}", linea[0]["text"]):
        return linea[0]["text"]
    partes = []
    for w in linea:
        alfanum = [c for c in w["text"] if c.isalnum()]
        digitos = sum(c.isdigit() for c in alfanum)
        if es_numero(w["text"]) or (digitos >= 3 and digitos * 2 >= len(alfanum)):
            break  # primera cifra del renglon (aunque el OCR la haya danado)
        partes.append(w["text"])
    return " ".join(partes)


def anos_en(linea):
    return [w for w in linea if re.fullmatch(r"(19|20)\d{2}", w["text"])]


@functools.lru_cache(maxsize=None)
def lineas_pdf(ruta, pagina_pdf):
    import pdfplumber

    with pdfplumber.open(ruta) as pdf:
        return lineas_de_pagina(pdf.pages[pagina_pdf - 1])


def releer_pdf(ruta, pagina_pdf, cuadro, celda):
    m = re.match(r"renglon: (.*?) \| columna: (.*?)(?: \| n: (\d+))?$", celda)
    if not m:
        return None, "celda sin formato reconocible"
    renglon, columna, n = m.group(1), m.group(2), m.group(3)
    lineas = lineas_pdf(str(ruta), int(pagina_pdf))
    texto = [" ".join(w["text"] for w in l) for l in lineas]

    # region del cuadro: desde su titulo hasta el titulo de otro cuadro
    mc = re.search(r"(\d+(?:\.\d+)*)\s*$", cuadro)
    num_cuadro = mc.group(1) if mc else None
    titulo = re.compile(r"Cuadro\s*(?:No\.?\s*)?(\d+(?:\.\d+)*)")
    inicio, fin = 0, len(lineas)
    for i, t in enumerate(texto):
        mt = titulo.search(t)
        if mt and mt.group(1) == num_cuadro:
            inicio = i
            break
    for i in range(inicio + 1, len(lineas)):
        mt = titulo.search(texto[i])
        if mt and mt.group(1) != num_cuadro:
            fin = i
            break

    objetivo = normaliza(renglon)

    def coincide_etiqueta(linea):
        e = normaliza(etiqueta(linea))
        if not e:
            return False
        if re.fullmatch(r"\d{4}", objetivo):
            return e == objetivo
        return e == objetivo or difflib.SequenceMatcher(None, e, objetivo).ratio() >= 0.75

    if n:  # columna por posicion ordinal de la cifra dentro del renglon
        for i in range(inicio, fin):
            if coincide_etiqueta(lineas[i]):
                celdas = celdas_numericas(lineas[i])
                if re.fullmatch(r"\d{4}", objetivo) and celdas and celdas[0]["text"] == renglon:
                    celdas = celdas[1:]
                k = int(n) - 1
                if k < len(celdas):
                    return a_decimal(celdas[k]["text"]), celdas[k]["text"]
                return None, "renglon sin esa columna"
        return None, "renglon no encontrado"

    # columna por ano: encabezado con varios anos y cifra alineada bajo el ano
    for i in range(inicio, fin):
        anos = anos_en(lineas[i])
        if len(anos) >= 3 and any(w["text"] == columna for w in anos):
            cab = next(w for w in anos if w["text"] == columna)
            centros = sorted((w["x0"] + w["x1"]) / 2 for w in anos)
            paso = min(b - a for a, b in zip(centros, centros[1:])) if len(centros) > 1 else 50
            cx = (cab["x0"] + cab["x1"]) / 2
            for j in range(i + 1, fin):
                if len(anos_en(lineas[j])) >= 3:
                    break
                if coincide_etiqueta(lineas[j]):
                    celdas = celdas_numericas(lineas[j])
                    if not celdas:
                        continue
                    mejor = min(celdas, key=lambda c: abs((c["x0"] + c["x1"]) / 2 - cx))
                    if abs((mejor["x0"] + mejor["x1"]) / 2 - cx) <= paso / 2:
                        return a_decimal(mejor["text"]), mejor["text"]
                    return None, "sin cifra bajo el ano"
            return None, "renglon no encontrado bajo el encabezado"
    return None, "encabezado con el ano no encontrado"


@functools.lru_cache(maxsize=None)
def datos_json(ruta):
    return json.loads(Path(ruta).read_text(encoding="utf-8"))[1]


def releer_json(ruta, celda):
    """celda = 'registro: <indice> | campo: value' dentro de la lista de datos (posicion 1)."""
    m = re.match(r"registro: (\d+) \| campo: (\w+)$", celda)
    datos = datos_json(str(ruta))
    v = datos[int(m.group(1))][m.group(2)]
    return (None if v is None else Decimal(repr(v))), repr(v)


@functools.lru_cache(maxsize=None)
def libro_xlsx(ruta):
    import openpyxl

    return openpyxl.load_workbook(ruta, data_only=True)


def releer_xlsx(ruta, hoja, celda):
    v = libro_xlsx(str(ruta))[hoja][celda].value
    return (None if v is None else Decimal(repr(v))), repr(v)


def paso_c():
    salida = []
    for fila in leer_csv(CSV_CIFRAS):
        ruta = FUENTES / fila["archivo"]
        registrado = Decimal(fila["valor"])
        if fila["lectura"] == "imagen":
            releido, coincide, nota = "", "no_aplica", "cifra leida de la imagen de la pagina"
        elif not ruta.exists():
            releido, coincide, nota = "", "no", "falta el archivo"
        else:
            try:
                sufijo = ruta.suffix.lower()
                if sufijo == ".pdf":
                    valor, crudo = releer_pdf(ruta, fila["pagina_pdf"], fila["cuadro"], fila["celda"])
                elif sufijo == ".xlsx":
                    valor, crudo = releer_xlsx(ruta, fila["hoja"], fila["celda"])
                else:
                    valor, crudo = releer_json(ruta, fila["celda"])
            except Exception as e:
                valor, crudo = None, f"error: {e}"
            if valor is None:
                releido, coincide, nota = "", "no", crudo
            else:
                releido = str(valor)
                coincide = "si" if valor == registrado else "no"
                nota = f"texto leido: {crudo}"
        salida.append({"anio": fila["anio"], "variable": fila["variable"], "archivo": fila["archivo"],
                       "valor_registrado": fila["valor"], "valor_releido": releido,
                       "coincide": coincide, "nota": nota})
    with open(CSV_VALIDACION, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["anio", "variable", "archivo", "valor_registrado",
                                          "valor_releido", "coincide", "nota"])
        w.writeheader()
        w.writerows(salida)
    conteo = {}
    for s in salida:
        conteo[s["coincide"]] = conteo.get(s["coincide"], 0) + 1
    print(f"(c) {len(salida)} cifras releidas: {conteo}")


# ---------------------------------------------------------------- (d) Penn World Table

def celda_texto(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(v)
    return str(v)


def paso_d():
    if not PWT_XLSX.exists():
        print("(d) falta fuentes/pwt1001.xlsx: no se rehacen los archivos de la Penn World Table")
        return
    import openpyxl

    libro = openpyxl.load_workbook(PWT_XLSX, read_only=True, data_only=True)
    hojas = {h.lower(): h for h in libro.sheetnames}
    hoja_datos = libro[hojas["data"]]
    filas = hoja_datos.iter_rows(values_only=True)
    encabezado = list(next(filas))
    i_pais = encabezado.index("countrycode")
    with open(CSV_PWT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(encabezado)
        n = 0
        for fila in filas:
            if fila[i_pais] == "MEX":
                w.writerow([celda_texto(v) for v in fila])
                n += 1
    print(f"(d) {n} filas de MEX en {CSV_PWT.name}")
    with open(CSV_LEYENDA, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for fila in libro[hojas["legend"]].iter_rows(values_only=True):
            w.writerow([celda_texto(v) for v in fila])
    print(f"(d) hoja Legend copiada en {CSV_LEYENDA.name}")


if __name__ == "__main__":
    paso_a()
    paso_b()
    paso_c()
    paso_d()
    sys.exit(0)
